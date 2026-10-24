#!/usr/bin/env python
"""
tests/test_core.py
Core modules: config, model loading, database wrapper, attendance engine, exporter.
Plain script, no pytest needed.

Checks:
1. config.json creation with defaults when missing.
2. model_loader: missing Trainer.yml handled; missing emotion model raises clear error.
3. Mongo: insert/query round-trip & weekly summary (or graceful disconnection).
4. Firebase with blank args: every method is a no-op.
5. Engine: cooldown, shared unknown cooldown, Present/Late, auto_attend off, manual_mark.
6. Exporter: both files created with identical column order.
"""

import csv
import json
import os
import sys
import tempfile
import hashlib
from datetime import datetime, timedelta

# Ensure project root is on sys.path
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

def get_real_cfg_hash():
    p = os.path.join(PROJECT_ROOT, "config.json")
    if os.path.exists(p):
        with open(p, "rb") as f:
            return hashlib.sha256(f.read()).hexdigest()
    return None

_cfg_hash_before = get_real_cfg_hash()

passed = 0
failed = 0


def check(label, condition, detail=""):
    global passed, failed
    if condition:
        passed += 1
        print(f"  [PASS] {label}")
    else:
        failed += 1
        print(f"  [FAIL] {label}  -- {detail}")


def section(title):
    print(f"\n{'=' * 60}")
    print(f"  {title}")
    print(f"{'=' * 60}")


# -------------------------------------------------------------------------
# 1. config.json creation
# -------------------------------------------------------------------------
section("1. config.json -- create with defaults when missing")

from core.model_loader import load_config, save_config, DEFAULT_CONFIG

with tempfile.TemporaryDirectory() as tmpdir:
    tmp_cfg_path = os.path.join(tmpdir, "config.json")
    cfg = load_config(tmp_cfg_path)
    check("File created on disk", os.path.exists(tmp_cfg_path))
    check("Has correct camera default", cfg.get("camera") == 0)
    check("Has emotions list", isinstance(cfg.get("emotions"), list) and len(cfg["emotions"]) == 7)
    check("Has attendance sub-config", "cooldown_minutes" in cfg.get("attendance", {}))

    # Round-trip: modify and reload
    cfg["camera"] = 2
    save_config(cfg, tmp_cfg_path)
    cfg2 = load_config(tmp_cfg_path)
    check("Save/reload preserves changes", cfg2["camera"] == 2)

# -------------------------------------------------------------------------
# 2. model_loader: recognizer + emotion model
# -------------------------------------------------------------------------
section("2. model_loader -- Trainer.yml / emotion model handling")

from core.model_loader import load_recognizer, load_emotion_model, load_haar_cascade

# Haar cascade
try:
    cascade = load_haar_cascade()
    check("Haar cascade loaded", not cascade.empty())
except Exception as e:
    check("Haar cascade loaded", False, str(e))

# Missing Trainer.yml -> non-fatal
recognizer, ready = load_recognizer("nonexistent_Trainer.yml")
check("Missing Trainer.yml is non-fatal", True)
check("recognizer_ready is False", ready is False)

# Missing emotion model -> FileNotFoundError with path in message
try:
    load_emotion_model("nonexistent_emotion_model.h5")
    check("Missing emotion model raises error", False, "no exception raised")
except FileNotFoundError as e:
    msg = str(e)
    check("Missing emotion model raises FileNotFoundError", True)
    check("Error message names the path", "nonexistent_emotion_model" in msg, msg)
except Exception as e:
    check("Missing emotion model raises FileNotFoundError", False, f"wrong exception: {type(e).__name__}: {e}")

# -------------------------------------------------------------------------
# 3. MongoDB
# -------------------------------------------------------------------------
section("3. MongoDB -- insert/query or graceful disconnection")

from db.mongo_client import MongoAttendanceDB

# Try connecting to local Mongo
mongo = MongoAttendanceDB("mongodb://localhost:27017", "face_attendance_test")

if mongo.is_connected:
    print("  (local MongoDB IS running -- testing insert/query round-trip)")

    # Clean up test data first
    mongo._collection.delete_many({"name": "__test__"})

    today_str = datetime.now().strftime("%Y-%m-%d")
    rec = {
        "person_id": 999,
        "name": "__test__",
        "date": today_str,
        "time_in": "08:00:00",
        "status": "Present",
        "emotion": "Happy",
        "confidence": 95.5,
        "manual": False,
    }
    result = mongo.insert_log(dict(rec))  # copy to avoid mutation issues
    check("insert_log returns record with id", result is not None and "id" in result)

    logs = mongo.get_logs_by_date(today_str)
    found = any(r["name"] == "__test__" for r in logs)
    check("get_logs_by_date finds the record", found)

    today_logs = mongo.get_today_logs()
    found_today = any(r["name"] == "__test__" for r in today_logs)
    check("get_today_logs finds the record", found_today)

    summary = mongo.get_weekly_summary()
    check("get_weekly_summary returns 7 entries", len(summary) == 7, f"got {len(summary)}")
    check("Each entry has required keys",
          all("date" in s and "present_count" in s and "late_count" in s and "unknown_count" in s for s in summary))

    emotions = mongo.get_today_emotion_counts()
    check("get_today_emotion_counts returns dict", isinstance(emotions, dict))

    # Clean up
    mongo._client.drop_database("face_attendance_test")
    mongo.close()
else:
    print("  (local MongoDB is NOT running -- testing graceful failure mode)")
    check("is_connected is False", mongo.is_connected is False)
    check("get_logs_by_date returns []", mongo.get_logs_by_date("2025-01-01") == [])
    check("get_today_logs returns []", mongo.get_today_logs() == [])
    check("get_weekly_summary returns 7 entries", len(mongo.get_weekly_summary()) == 7)
    check("get_today_emotion_counts returns {}", mongo.get_today_emotion_counts() == {})
    check("insert_log returns None", mongo.insert_log({"test": True}) is None)

# -------------------------------------------------------------------------
# 4. Firebase with blank args
# -------------------------------------------------------------------------
section("4. Firebase -- blank args -> silent no-op")

from db.firebase_client import FirebaseClient

fb = FirebaseClient("", "")
check("is_enabled is False", fb.is_enabled is False)
# push_attendance must not raise
try:
    fb.push_attendance({"person_id": 1, "name": "test"})
    check("push_attendance is a no-op (no raise)", True)
except Exception as e:
    check("push_attendance is a no-op (no raise)", False, str(e))

fb2 = FirebaseClient("", "https://something.firebaseio.com")
check("One blank arg still disables", fb2.is_enabled is False)

fb3 = FirebaseClient("some_path.json", "")
check("Other blank arg still disables", fb3.is_enabled is False)

# -------------------------------------------------------------------------
# 5. AttendanceEngine (with fake in-memory DB)
# -------------------------------------------------------------------------
section("5. AttendanceEngine -- cooldown, status, auto_attend, manual")


class FakeDB:
    """Minimal in-memory stand-in for MongoAttendanceDB."""
    def __init__(self):
        self.logs = []

    @property
    def is_connected(self):
        return True

    def insert_log(self, record):
        record["id"] = str(len(self.logs) + 1)
        record["created_at"] = datetime.utcnow()
        self.logs.append(record)
        return record

    def get_today_logs(self):
        today = datetime.now().strftime("%Y-%m-%d")
        return [r for r in self.logs if r.get("date") == today]


class FakeFirebase:
    """Silent stand-in for FirebaseClient."""
    is_enabled = False
    def push_attendance(self, record):
        pass


from attendance.engine import AttendanceEngine

fake_db = FakeDB()
fake_fb = FakeFirebase()

# 5a. Cooldown blocks repeat within window
current_time = datetime(2025, 6, 15, 8, 0, 0)
engine = AttendanceEngine(
    {
        "names": {1: "Alice"},
        "attendance": {
            "cooldown_minutes": 60,
            "late_after": "09:30",
            "working_hours_end": "17:30",
            "auto_attend": True,
        },
    },
    fake_db,
    fake_fb,
    clock=lambda: current_time,
)

r1 = engine.process_detection(1, True, "Happy", 95.0)
check("First detection logs", r1 is not None)

r2 = engine.process_detection(1, True, "Happy", 95.0)
check("Repeat within cooldown blocked", r2 is None)

# Advance time past cooldown
current_time = datetime(2025, 6, 15, 9, 5, 0)
r3 = engine.process_detection(1, True, "Happy", 95.0)
check("After cooldown, logs again", r3 is not None)

# 5b. Unknown faces share one cooldown
current_time = datetime(2025, 6, 15, 10, 0, 0)
r_u1 = engine.process_detection(-1, False, "Sad", 30.0)
check("First unknown face logs", r_u1 is not None)
check("Unknown status is 'Unknown'", r_u1["status"] == "Unknown")

r_u2 = engine.process_detection(-2, False, "Angry", 20.0)
check("Second unknown within cooldown blocked (shared)", r_u2 is None)

# 5c. Before/after late_after -> Present / Late (on a working day: 16 June 2025 is a Monday)
fake_db2 = FakeDB()
time_early = datetime(2025, 6, 16, 8, 30, 0)
engine2 = AttendanceEngine(
    {"names": {2: "Bob"}, "attendance": {"cooldown_minutes": 0, "late_after": "09:30",
                                          "working_hours_end": "17:30", "auto_attend": True}},
    fake_db2, fake_fb, clock=lambda: time_early,
)
r_present = engine2.process_detection(2, True, "Neutral", 80.0)
check("Before late_after -> Present", r_present is not None and r_present["status"] == "Present")

time_late = datetime(2025, 6, 16, 10, 0, 0)
engine2._clock = lambda: time_late
r_late = engine2.process_detection(3, True, "Happy", 85.0)
check("After late_after -> Late", r_late is not None and r_late["status"] == "Late")

# 5d. auto_attend off -> no-op
engine2.set_auto_attend(False)
r_off = engine2.process_detection(4, True, "Happy", 90.0)
check("auto_attend off -> returns None", r_off is None)

# 5e. manual_mark bypasses cooldown
engine2.set_auto_attend(True)
# Person 2 was already logged, but manual_mark should still work
m1 = engine2.manual_mark(2, "Bob", "Present")
check("manual_mark returns record", m1 is not None)
check("manual_mark record has manual=True", m1["manual"] is True)

m2 = engine2.manual_mark(2, "Bob", "Late")
check("manual_mark ignores cooldown (second call works)", m2 is not None)

# 5f. get_today_stats with mixed Present, Late, and Unknown records
fake_db_stats = FakeDB()
engine_stats = AttendanceEngine(
    {"names": {1: "Alice", 2: "Bob"}, "attendance": {"cooldown_minutes": 0, "late_after": "09:30",
                                                      "working_hours_end": "17:30", "auto_attend": True}},
    fake_db_stats, fake_fb, clock=lambda: datetime.now(),
)
# Case 1: Only Unknown records
today_str = datetime.now().strftime("%Y-%m-%d")
fake_db_stats.logs = [
    {"person_id": -1, "name": "Unknown", "date": today_str, "status": "Unknown"}
    for _ in range(6)
]
stats_unk = engine_stats.get_today_stats()
check("All Unknown records: present is 0", stats_unk["present"] == 0)
check("All Unknown records: late is 0", stats_unk["late"] == 0)
check("All Unknown records: total is 6 (includes all records)", stats_unk["total"] == 6)
check("All Unknown records: present + late is 0 (excludes unknown)", stats_unk["present"] + stats_unk["late"] == 0)

# Case 2: Mixed Present, Late, and Unknown records
fake_db_stats.logs = [
    {"person_id": 1, "name": "Alice", "date": today_str, "status": "Present"},
    {"person_id": 2, "name": "Bob", "date": today_str, "status": "Present"},
    {"person_id": 3, "name": "Charlie", "date": today_str, "status": "Late"},
    {"person_id": -1, "name": "Unknown", "date": today_str, "status": "Unknown"},
    {"person_id": -1, "name": "Unknown", "date": today_str, "status": "Unknown"},
]
stats_mix = engine_stats.get_today_stats()
check("Mixed records: present is 2", stats_mix["present"] == 2)
check("Mixed records: late is 1", stats_mix["late"] == 1)
check("Mixed records: total is 5 (all records including unknown)", stats_mix["total"] == 5)
check("Mixed records: present + late is 3 (excludes unknown)", stats_mix["present"] + stats_mix["late"] == 3)

# Case 3: Mixed set of 1 Present, 1 Late, 1 Unknown
fake_db_stats.logs = [
    {"person_id": 1, "name": "Alice", "date": today_str, "status": "Present"},
    {"person_id": 2, "name": "Bob", "date": today_str, "status": "Late"},
    {"person_id": -1, "name": "Unknown", "date": today_str, "status": "Unknown"},
]
stats_111 = engine_stats.get_today_stats()
check("1 Present, 1 Late, 1 Unknown: total is 3", stats_111["total"] == 3)
check("1 Present, 1 Late, 1 Unknown: present is 1", stats_111["present"] == 1)
check("1 Present, 1 Late, 1 Unknown: late is 1", stats_111["late"] == 1)
check("1 Present, 1 Late, 1 Unknown: present + late is 2 (excludes unknown)", stats_111["present"] + stats_111["late"] == 2)

# -------------------------------------------------------------------------
# 6. Exporter -- both files, identical column order
# -------------------------------------------------------------------------
section("6. Exporter -- Excel & CSV with identical columns")

from attendance.exporter import export_to_excel, export_to_csv, COLUMNS

sample = [
    {"name": "Alice", "person_id": 1, "time_in": "08:00:00", "status": "Present",
     "emotion": "Happy", "confidence": 95.5, "manual": False},
    {"name": "Bob", "person_id": 2, "time_in": "10:00:00", "status": "Late",
     "emotion": "Neutral", "confidence": 80.0, "manual": True},
]

with tempfile.TemporaryDirectory() as tmpdir:
    xlsx_path = os.path.join(tmpdir, "test.xlsx")
    csv_path = os.path.join(tmpdir, "test.csv")

    export_to_excel(sample, xlsx_path)
    export_to_csv(sample, csv_path)

    check("Excel file created", os.path.exists(xlsx_path))
    check("CSV file created", os.path.exists(csv_path))

    # Read CSV header
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        csv_header = next(reader)
    check("CSV columns match expected order", csv_header == COLUMNS, f"got {csv_header}")

    # Read Excel header
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path)
    ws = wb.active
    xlsx_header = [cell.value for cell in ws[1]]
    check("Excel columns match expected order", xlsx_header == COLUMNS, f"got {xlsx_header}")

    check("CSV and Excel columns identical", csv_header == xlsx_header)

    # Check data rows
    csv_rows = list(csv.reader(open(csv_path, "r", encoding="utf-8")))
    check("CSV has header + 2 data rows", len(csv_rows) == 3, f"got {len(csv_rows)} rows")
    check("Excel has header + 2 data rows", ws.max_row == 3, f"got {ws.max_row} rows")

# -------------------------------------------------------------------------
# 7. Config integrity check
# -------------------------------------------------------------------------
section("7. Real config.json integrity")
_cfg_hash_after = get_real_cfg_hash()
check("Real config.json unchanged during test run", _cfg_hash_before == _cfg_hash_after)

# -------------------------------------------------------------------------
# Summary
# -------------------------------------------------------------------------
print(f"\n{'=' * 60}")
print(f"  RESULTS: {passed} passed, {failed} failed, {passed + failed} total")
print(f"{'=' * 60}")
sys.exit(1 if failed else 0)
