"""
tests/test_attendance_dashboard.py
Attendance and Dashboard end to end: routes, database operations, exports,
snapshot serving, and behaviour when the database is unreachable.
"""

import csv
import hashlib
import io
import json
import os
import shutil
import sys
import tempfile
import time
from datetime import datetime, timedelta

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

def get_real_cfg_hash():
    p = os.path.join(PROJECT_ROOT, "config.json")
    if os.path.exists(p):
        with open(p, "rb") as f:
            return hashlib.sha256(f.read()).hexdigest()
    return None

_cfg_hash_before = get_real_cfg_hash()

# Setup isolated temp config
tmp_dir = tempfile.mkdtemp(prefix="face_web_test_")
tmp_cfg_path = os.path.join(tmp_dir, "config.json")
real_cfg_path = os.path.join(PROJECT_ROOT, "config.json")
if os.path.exists(real_cfg_path):
    with open(real_cfg_path, "r", encoding="utf-8") as f:
        cfg_copy = json.load(f)
else:
    from core.model_loader import DEFAULT_CONFIG
    cfg_copy = dict(DEFAULT_CONFIG)

cfg_copy["mongodb_db"] = "face_attendance_test"
with open(tmp_cfg_path, "w", encoding="utf-8") as f:
    json.dump(cfg_copy, f, indent=2)

os.environ["FACE_WEB_CONFIG"] = tmp_cfg_path

import cv2
import numpy as np
from openpyxl import load_workbook

import server
from attendance.engine import AttendanceEngine
from attendance.exporter import export_to_csv, export_to_excel
from db.mongo_client import MongoAttendanceDB

passed = 0
failed = 0


def check(label, condition, detail=""):
    global passed, failed
    if condition:
        passed += 1
        print(f"  [PASS] {label}")
    else:
        failed += 1
        print(f"  [FAIL] {label}  -- {detail}")


def section(title):
    print(f"\n{'=' * 60}")
    print(f"  {title}")
    print(f"{'=' * 60}")


# Switch DB to face_attendance_test
test_db = MongoAttendanceDB("mongodb://localhost:27017", "face_attendance_test")
server.mongo_db = test_db
server.engine._db = test_db

if test_db.is_connected:
    test_db._collection.delete_many({})

from helpers import signed_in_client  # noqa: E402
client = signed_in_client(server)

# ── 1. Seed records & Verify Listing / Counts ────────────────────────
section("1. Seed records (Present, Late, Unknown) & Verify Listing / Counts")

today = datetime.now().strftime("%Y-%m-%d")
yesterday = (datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d")

# Manual marks are for registered people only
from helpers import register_people  # noqa: E402
register_people(server, {101: "Alice Cooper", 102: "Bob Marley"})

# Manual Mark: Present
r_pres = client.post("/api/attendance/manual", json={"person_id": 101, "name": "Alice Cooper", "status": "Present"})
check("Manual mark Present status 200", r_pres.status_code == 200)

# Manual Mark: Late
r_late = client.post("/api/attendance/manual", json={"person_id": 102, "name": "Bob Marley", "status": "Late"})
check("Manual mark Late status 200", r_late.status_code == 200)

# Unknown face seen by the camera (manual marks cannot be Unknown)
r_unk = server.engine.process_detection(-1, False, "", 0.0)
check("Unknown face recorded", r_unk is not None)

# Detection with face_crop: Present
fake_crop = np.zeros((100, 100), dtype=np.uint8)
det_rec = server.engine.process_detection(103, True, "Happy", 35.5, face_crop=fake_crop)
check("Engine process_detection logged", det_rec is not None)
check("det_rec has_snapshot is True", det_rec.get("has_snapshot") is True)
det_id = det_rec.get("id")

# Fetch records for today
res_today = client.get(f"/api/attendance?date={today}")
today_records = res_today.get_json()
check("Today has 4 records", len(today_records) == 4, f"got {len(today_records)}")

# Verify filter counts
count_all = len(today_records)
count_present = sum(1 for r in today_records if r["status"] == "Present")
count_late = sum(1 for r in today_records if r["status"] == "Late")
count_unknown = sum(1 for r in today_records if r["status"] == "Unknown")

expected_late = 2 if det_rec.get("status") == "Late" else 1
expected_present = 1 if det_rec.get("status") == "Late" else 2

check("Filter count All == 4", count_all == 4)
check("Filter count Present matches expected", count_present == expected_present)
check("Filter count Late matches expected", count_late == expected_late)
check("Filter count Unknown == 1", count_unknown == 1)
check("Sum of filters matches Total", count_present + count_late + count_unknown == count_all)

# ── Live Stats Present check: get_today_stats reflects known attendees (Present + Late), never Unknown
stats_today = server.engine.get_today_stats()
check("get_today_stats present matches expected", stats_today["present"] == expected_present)
check("get_today_stats late matches expected", stats_today["late"] == expected_late)
check("get_today_stats total equals all records including unknown", stats_today["total"] == count_all)
check("get_today_stats total is 4", stats_today["total"] == 4)
check("get_today_stats present + late excludes Unknown (3)", stats_today["present"] + stats_today["late"] == 3)

# ── 2. Snapshot Verification (Detected vs Manual) ────────────────────
section("2. Snapshot Verification (Detected vs Manual)")

# Detected record has snapshot
res_snap = client.get(f"/api/attendance/{det_id}/snapshot")
check("Detected record snapshot returns 200", res_snap.status_code == 200)
check("Snapshot Content-Type is image/jpeg", res_snap.headers.get("Content-Type") == "image/jpeg")
res_snap.close()

# Clean up created snapshot file
snap_file = os.path.join(PROJECT_ROOT, "snapshots", f"{det_id}.jpg")
if os.path.exists(snap_file):
    os.remove(snap_file)

# Manual record does NOT have snapshot
manual_rec_id = r_pres.get_json()["record"]["id"]
res_manual_snap = client.get(f"/api/attendance/{manual_rec_id}/snapshot")
check("Manual record snapshot returns 404", res_manual_snap.status_code == 404)

# ── 3. Date Navigation & Past Date Records ───────────────────────────
section("3. Date Navigation & Past Date Records")

# Insert record for yesterday directly
past_rec = {
    "person_id": 104,
    "name": "Past Visitor",
    "date": yesterday,
    "time_in": "10:15:00",
    "status": "Late",
    "emotion": "Neutral",
    "confidence": 55.0,
    "manual": True,
    "has_snapshot": False,
}
test_db.insert_log(past_rec)

res_past = client.get(f"/api/attendance?date={yesterday}")
past_records = res_past.get_json()
check("Past date returns records", len(past_records) == 1)
check("Past record name matches", past_records[0]["name"] == "Past Visitor")
check("Past record date matches", past_records[0]["date"] == yesterday)

# ── 4. Excel & CSV Export Verification ───────────────────────────────
section("4. Excel & CSV Export Verification")

expected_cols = ["Date", "Name", "ID", "Time In", "Status", "Emotion", "Confidence", "Manual", "Review"]

# CSV Export for today
res_csv = client.get(f"/api/export/csv?date={today}")
check("CSV export returns 200", res_csv.status_code == 200)
check("CSV Content-Type is text/csv", "text/csv" in res_csv.headers.get("Content-Type", ""))
check("CSV filename in header", f"attendance_{today}.csv" in res_csv.headers.get("Content-Disposition", ""))

csv_rows = list(csv.reader(io.StringIO(res_csv.data.decode("utf-8"))))
check("CSV header matches canonical order", csv_rows[0] == expected_cols, str(csv_rows[0]))
check("CSV data rows count matches today records (4)", len(csv_rows) - 1 == 4, f"got {len(csv_rows) - 1}")

# Excel Export for today
res_xlsx = client.get(f"/api/export/excel?date={today}")
check("Excel export returns 200", res_xlsx.status_code == 200)
check("Excel filename in header", f"attendance_{today}.xlsx" in res_xlsx.headers.get("Content-Disposition", ""))

wb = load_workbook(io.BytesIO(res_xlsx.data))
ws = wb.active
xlsx_header = [cell.value for cell in ws[1]]
check("Excel header matches canonical order", xlsx_header == expected_cols, str(xlsx_header))
check("Excel data rows count matches today records (4)", ws.max_row - 1 == 4, f"got {ws.max_row - 1}")

# ── 5. Manual Mark Validation (Invalid Status & Bad Inputs) ──────────
section("5. Manual Mark Validation (Invalid Status & Bad Inputs)")

# Invalid status
r_bad_stat = client.post("/api/attendance/manual", json={"person_id": 1, "name": "Bob", "status": "InvalidStatus"})
check("Invalid status returns 400", r_bad_stat.status_code == 400)
check("Error message indicates status issue", "Status must be Present or Late" in r_bad_stat.get_json().get("error", ""))

# Missing name
r_no_name = client.post("/api/attendance/manual", json={"person_id": 1, "status": "Present"})
check("Missing name returns 400", r_no_name.status_code == 400)

# Non-integer person_id
r_str_id = client.post("/api/attendance/manual", json={"person_id": "one", "name": "Bob", "status": "Present"})
check("String person_id returns 400", r_str_id.status_code == 400)

r_bool_id = client.post("/api/attendance/manual", json={"person_id": False, "name": "Bob", "status": "Present"})
check("Boolean person_id returns 400", r_bool_id.status_code == 400)

# ── 6. Dashboard Stats APIs Verification ─────────────────────────────
section("6. Dashboard Stats APIs Verification")

# /api/stats/weekly
res_w = client.get("/api/stats/weekly")
check("Weekly stats returns 200", res_w.status_code == 200)
w_list = res_w.get_json()
check("Weekly stats has 7 entries", len(w_list) == 7)
today_entry = next((item for item in w_list if item["date"] == today), None)
check("Today entry present in weekly summary", today_entry is not None)
if today_entry:
    check("Today present count matches", today_entry["present_count"] == expected_present)
    check("Today late count matches", today_entry["late_count"] == expected_late)
    check("Today unknown count matches", today_entry["unknown_count"] == 1)

# /api/stats/emotions
res_e = client.get("/api/stats/emotions")
check("Emotions stats returns 200", res_e.status_code == 200)
e_dict = res_e.get_json()
check("Happy emotion counted", e_dict.get("Happy", 0) >= 1)

# /api/stats/persons
res_p = client.get("/api/stats/persons")
check("Persons stats returns 200", res_p.status_code == 200)
p_list = res_p.get_json()
check("Persons stats has entries", len(p_list) >= 3)
p_names = [p["name"] for p in p_list]
check("Alice Cooper in persons list", "Alice Cooper" in p_names)
check("Bob Marley in persons list", "Bob Marley" in p_names)

# ── Test: A day with only manual records gives empty emotion count and no doughnut slice
if test_db.is_connected:
    test_db._collection.delete_many({"manual": False})
    res_man_e = client.get("/api/stats/emotions")
    check("Manual-only day /api/stats/emotions returns 200", res_man_e.status_code == 200)
    man_e_data = res_man_e.get_json()
    check("Day with only manual records gives empty emotion count", man_e_data == {})
    # Verify no doughnut slices in dashboard chart logic
    doughnut_entries = [
        (emo, cnt) for emo, cnt in man_e_data.items()
        if cnt > 0 and emo and emo != "--"
    ]
    check("Day with only manual records produces no doughnut slice", len(doughnut_entries) == 0)

# ── 7. Disconnected MongoDB Resilience ───────────────────────────────
section("7. Disconnected MongoDB Resilience (Empty States, No 500)")

offline_db = MongoAttendanceDB("mongodb://localhost:29999", "fake_db")
server.mongo_db = offline_db

res_off_att = client.get("/api/attendance")
check("/api/attendance returns 200 (not 500)", res_off_att.status_code == 200)
check("/api/attendance returns []", res_off_att.get_json() == [])

res_off_w = client.get("/api/stats/weekly")
check("/api/stats/weekly returns 200 (not 500)", res_off_w.status_code == 200)
check("/api/stats/weekly returns 7 zeroed entries", len(res_off_w.get_json()) == 7)

res_off_e = client.get("/api/stats/emotions")
check("/api/stats/emotions returns 200", res_off_e.status_code == 200)
check("/api/stats/emotions returns {}", res_off_e.get_json() == {})

res_off_p = client.get("/api/stats/persons")
check("/api/stats/persons returns 200", res_off_p.status_code == 200)
check("/api/stats/persons returns []", res_off_p.get_json() == [])

# Cleanup test db
if test_db.is_connected:
    test_db._client.drop_database("face_attendance_test")

shutil.rmtree(tmp_dir, ignore_errors=True)

# ── 8. Real config.json integrity ───────────────────────────────────
section("8. Real config.json integrity")
_cfg_hash_after = get_real_cfg_hash()
check("Real config.json unchanged during test run", _cfg_hash_before == _cfg_hash_after)

# ── Summary ──────────────────────────────────────────────────────────
print(f"\n{'=' * 60}")
print(f"  RESULTS: {passed} passed, {failed} failed, {passed + failed} total")
print(f"{'=' * 60}")
sys.exit(1 if failed else 0)
